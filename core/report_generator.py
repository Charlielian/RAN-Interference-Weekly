# -*- coding: utf-8 -*-
"""
干扰小区周报生成模块
集成《干扰小区周报_脚本_融合版.py》的核心算法：
包含：干扰电平判断、干扰小区标识判断双Sheet，阳江统计、区域分布、全省21地市三类排名。
"""
import os
import re
import gc
from datetime import datetime, timedelta
from collections import defaultdict

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.config import (
    CITY_CLASS, ALL_CITIES, SOURCE_4G_DIR, SOURCE_5G_DIR, REPORT_DIR
)

# 干扰门限 (dBm)
THRESHOLD_4G     = -105
THRESHOLD_5G_700 = -105
THRESHOLD_5G_OTH = -107

# 5G 频段映射
BAND_5G_MAP = {
    "700M":   "700M",
    "2.6GHz": "2.6G",
    "4.9GHz": "4.9G",
}

# 4G 频段映射
BAND_4G_MAP = {
    "FD":    "FDD1800",
    "FG":    "FDD900",
    "LF":    "F频",
    "LD":    "D频",
    "LE":    "E频",
    "A频段": "A频段",
    "NB":    "NB-IoT",
}

def extract_yj_area(cell_name):
    """
    从阳江小区名称中提取区域 (阳西/江城/南区/阳东/阳春)。
    兼容形如：
    - 阳江-GZ-阳江阳西...
    - CBN-阳江江城南区... (优先划入南区)
    - CBN-阳江海陵区... (海陵属于南区)
    - CBN-阳江阳东... / 阳西... / 阳春...
    """
    if not isinstance(cell_name, str):
        return ""
    if "海陵" in cell_name or "阳江南区" in cell_name or "江城南区" in cell_name:
        return "南区"
    if "阳江阳西" in cell_name:
        return "阳西"
    if "阳江阳东" in cell_name:
        return "阳东"
    if "阳江阳春" in cell_name:
        return "阳春"
    if "阳江江城" in cell_name:
        return "江城"
    if "南区" in cell_name:
        return "南区"
    if "阳西" in cell_name:
        return "阳西"
    if "阳东" in cell_name:
        return "阳东"
    if "阳春" in cell_name:
        return "阳春"
    if "江城" in cell_name:
        return "江城"
    return ""

def map_band_4g(b):
    return BAND_4G_MAP.get(b, b or "")

def map_band_5g(b):
    return BAND_5G_MAP.get(b, b or "")

def threshold_for(system, band):
    if system == "4G":
        return THRESHOLD_4G
    return THRESHOLD_5G_700 if band == "700M" else THRESHOLD_5G_OTH


# 样式常量
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
SUB_FILL  = PatternFill(start_color="8FAADC", end_color="8FAADC", fill_type="solid")
HEAD_FONT = Font(bold=True, color="FFFFFF")

def style_header(cell, text=None):
    if text is not None:
        cell.value = text
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER

def style_sub_header(cell, text=None):
    if text is not None:
        cell.value = text
    cell.font = Font(bold=True, color="000000")
    cell.fill = SUB_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def process_city_4g(city: str, day_dates: list, log_fn=None) -> tuple:
    all_data = []
    for date in day_dates:
        file_path = os.path.join(SOURCE_4G_DIR, f"4G干扰小区_{date.strftime('%Y-%m-%d')}_{city}.xlsx")
        if not os.path.exists(file_path):
            continue
        try:
            try:
                df = pd.read_excel(file_path, dtype={"CGI": str}, engine="calamine")
            except Exception:
                df = pd.read_excel(file_path, dtype={"CGI": str})
        except Exception as e:
            if log_fn:
                log_fn(f"[WARN] 读取4G失败跳过: {file_path} ({e})")
            continue
        
        need_cols = ["CGI", "小区名", "频段", "平均干扰电平", "是否干扰小区"]
        for c in need_cols:
            if c not in df.columns:
                df[c] = None
        df = df[need_cols].copy()
        df["日期"] = date
        df["CGI"] = df["CGI"].astype(str).str.strip()
        df = df[(df["CGI"] != "None") & (df["CGI"] != "") & (df["CGI"] != "nan")]
        all_data.append(df)
        del df
        gc.collect()

    if not all_data:
        return 0, [], []

    combined = pd.concat(all_data, ignore_index=True)
    del all_data
    gc.collect()

    agg = combined.groupby(["CGI", "日期"]).agg({
        "小区名": "first",
        "频段": "first",
        "平均干扰电平": "mean",
        "是否干扰小区": "first"
    }).reset_index()

    pivot = agg.pivot(index="CGI", columns="日期", values="平均干扰电平")
    pivot_interf = agg.pivot(index="CGI", columns="日期", values="是否干扰小区")
    meta = agg.groupby("CGI").agg({"小区名": "first", "频段": "first"})

    del combined, agg
    gc.collect()

    total = len(pivot)
    level_records = []
    flag_records = []

    for cgi in pivot.index:
        row_meta = meta.loc[cgi]
        band = str(row_meta["频段"]) if row_meta["频段"] else ""
        day_vals = {d: pivot.loc[cgi, d] if d in pivot.columns else None for d in day_dates}
        day_interf = {d: pivot_interf.loc[cgi, d] if d in pivot_interf.columns else None for d in day_dates}
        valid = [v for v in day_vals.values() if v is not None and not np.isnan(v)]
        avg = sum(valid) / len(valid) if valid else float("nan")
        thr = THRESHOLD_4G

        # 判断1：按干扰电平门限
        level_exceed_days = sum(1 for v in valid if v > thr)
        if level_exceed_days >= 3:
            record = {
                "制式": "4G", "地市": city, "CGI": cgi, "小区名": row_meta["小区名"],
                "频段": band, "门限": thr, "干扰天数": level_exceed_days, "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            level_records.append(record)

        # 判断2：按是否干扰小区标识
        flag_exceed_days = sum(1 for v in day_interf.values() if v is not None and (v == True or str(v).lower() == 'true' or str(v) == '1'))
        if flag_exceed_days >= 3:
            record = {
                "制式": "4G", "地市": city, "CGI": cgi, "小区名": row_meta["小区名"],
                "频段": band, "门限": thr, "干扰天数": flag_exceed_days, "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            flag_records.append(record)

    del pivot, pivot_interf, meta
    gc.collect()
    return total, level_records, flag_records


def process_city_5g(city: str, day_dates: list, log_fn=None) -> tuple:
    all_data = []
    for date in day_dates:
        file_path = os.path.join(SOURCE_5G_DIR, f"5G干扰小区_{date.strftime('%Y-%m-%d')}_{city}.xlsx")
        if not os.path.exists(file_path):
            continue
        try:
            try:
                df = pd.read_excel(file_path, dtype={"CGI": str}, engine="calamine")
            except Exception:
                df = pd.read_excel(file_path, dtype={"CGI": str})
        except Exception as e:
            if log_fn:
                log_fn(f"[WARN] 读取5G失败跳过: {file_path} ({e})")
            continue
        
        need_cols = ["CGI", "小区名", "频段", "全频段均值", "是否干扰小区"]
        for c in need_cols:
            if c not in df.columns:
                df[c] = None
        df = df[need_cols].copy()
        df["日期"] = date
        df["CGI"] = df["CGI"].astype(str).str.strip()
        df = df[(df["CGI"] != "None") & (df["CGI"] != "") & (df["CGI"] != "nan")]
        all_data.append(df)
        del df
        gc.collect()

    if not all_data:
        return 0, [], []

    combined = pd.concat(all_data, ignore_index=True)
    del all_data
    gc.collect()

    agg = combined.groupby(["CGI", "日期"]).agg({
        "小区名": "first",
        "频段": "first",
        "全频段均值": "mean",
        "是否干扰小区": "first"
    }).reset_index()

    pivot = agg.pivot(index="CGI", columns="日期", values="全频段均值")
    pivot_interf = agg.pivot(index="CGI", columns="日期", values="是否干扰小区")
    meta = agg.groupby("CGI").agg({"小区名": "first", "频段": "first"})

    del combined, agg
    gc.collect()

    total = len(pivot)
    level_records = []
    flag_records = []

    for cgi in pivot.index:
        row_meta = meta.loc[cgi]
        band = str(row_meta["频段"]) if row_meta["频段"] else ""
        thr = THRESHOLD_5G_700 if "700" in band.upper() else THRESHOLD_5G_OTH
        day_vals = {d: pivot.loc[cgi, d] if d in pivot.columns else None for d in day_dates}
        day_interf = {d: pivot_interf.loc[cgi, d] if d in pivot_interf.columns else None for d in day_dates}
        valid = [v for v in day_vals.values() if v is not None and not np.isnan(v)]
        avg = sum(valid) / len(valid) if valid else float("nan")

        # 判断1：按干扰电平门限
        level_exceed_days = sum(1 for v in valid if v > thr)
        if level_exceed_days >= 3:
            record = {
                "制式": "5G", "地市": city, "CGI": cgi, "小区名": row_meta["小区名"],
                "频段": band, "门限": thr, "干扰天数": level_exceed_days, "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            level_records.append(record)

        # 判断2：按是否干扰小区标识
        flag_exceed_days = sum(1 for v in day_interf.values() if v is not None and (v == True or str(v).lower() == 'true' or str(v) == '1'))
        if flag_exceed_days >= 3:
            record = {
                "制式": "5G", "地市": city, "CGI": cgi, "小区名": row_meta["小区名"],
                "频段": band, "门限": thr, "干扰天数": flag_exceed_days, "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            flag_records.append(record)

    del pivot, pivot_interf, meta
    gc.collect()
    return total, level_records, flag_records


def filter_yangjiang(records: list, system: str, log_fn=None) -> pd.DataFrame:
    df = pd.DataFrame(records)
    if df.empty:
        return df
    df = df[df["地市"] == "阳江"].copy()
    df["区域"] = df["小区名"].apply(extract_yj_area)
    if system == "4G":
        df["频段"] = df["频段"].apply(map_band_4g)
    else:
        df["频段"] = df["频段"].apply(map_band_5g)
    df = df.sort_values(by=["干扰天数", "平均干扰电平"], ascending=[False, False]).reset_index(drop=True)
    return df


def compute_province_ranking(rank_data: list) -> pd.DataFrame:
    df = pd.DataFrame(rank_data)
    if df.empty:
        return pd.DataFrame()
    df["干扰占比"] = df.apply(lambda r: r["干扰小区数"] / r["总小区数"] if r["总小区数"] else 0.0, axis=1)
    df["全省排名"] = df.groupby("制式")["干扰占比"].rank(method="min", ascending=True).astype(int)
    df["类别"] = df["地市"].map({c: k for k, v in CITY_CLASS.items() for c in v})
    df["三类排名"] = df.groupby(["制式", "类别"])["干扰占比"].rank(method="min", ascending=True).astype(int)
    return df[["类别", "地市", "制式", "总小区数", "干扰小区数", "干扰占比", "全省排名", "三类排名"]]


def write_main_table(ws, yj_4g: pd.DataFrame, yj_5g: pd.DataFrame, day_cols: list):
    headers = ["制式", "区域", "CGI", "小区名", "频段", "门限", "干扰天数", "平均干扰电平"] + day_cols
    for col_idx, h in enumerate(headers, 1):
        style_header(ws.cell(1, col_idx, h))

    r = 2
    for df, system in [(yj_4g, "4G"), (yj_5g, "5G")]:
        if df is None or df.empty:
            continue
        for _, row in df.iterrows():
            ws.cell(r, 1, system)
            ws.cell(r, 2, row.get("区域", ""))
            ws.cell(r, 3, str(row.get("CGI", "")))
            ws.cell(r, 4, row.get("小区名", ""))
            ws.cell(r, 5, row.get("频段", ""))
            ws.cell(r, 6, row.get("门限", ""))
            ws.cell(r, 7, int(row.get("干扰天数", 0)))
            val_avg = row.get("平均干扰电平")
            ws.cell(r, 8, round(float(val_avg), 2) if val_avg is not None and not np.isnan(val_avg) else "")
            for i, d in enumerate(day_cols):
                v = row.get(d)
                ws.cell(r, 9 + i, round(float(v), 2) if v is not None and not np.isnan(v) else "")
            for c in range(1, len(headers) + 1):
                ws.cell(r, c).border = BORDER
            r += 1

    widths = [6, 8, 22, 32, 10, 8, 10, 14] + [10] * len(day_cols)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_yj_summary_block(ws, yj_4g, yj_5g, yj_4g_total, yj_5g_total, prev_4g, prev_5g, start_col=22, start_row=4):
    c = start_col
    ws.merge_cells(start_row=start_row, start_column=c, end_row=start_row, end_column=c + 6)
    style_header(ws.cell(start_row, c, "阳江45G干扰小区"))
    sub_headers = ["制式", "总小区数", "本周干扰小区数", "本周干扰占比", "上周干扰占比", "环比增幅", "全省排名", "三类排名"]
    for i, h in enumerate(sub_headers):
        style_sub_header(ws.cell(start_row + 1, c + i, h))

    yj_4g_cnt = len(yj_4g) if yj_4g is not None and not yj_4g.empty else 0
    yj_5g_cnt = len(yj_5g) if yj_5g is not None and not yj_5g.empty else 0

    ratio_4g = yj_4g_cnt / yj_4g_total if yj_4g_total else 0.0
    ratio_5g = yj_5g_cnt / yj_5g_total if yj_5g_total else 0.0

    rows_data = [
        ("4G", yj_4g_total, yj_4g_cnt, ratio_4g, prev_4g),
        ("5G", yj_5g_total, yj_5g_cnt, ratio_5g, prev_5g),
    ]

    for ri, (sys_name, total, cnt, ratio, prev) in enumerate(rows_data):
        curr_r = start_row + 2 + ri
        ws.cell(curr_r, c, sys_name)
        ws.cell(curr_r, c + 1, total)
        ws.cell(curr_r, c + 2, cnt)
        ws.cell(curr_r, c + 3, round(ratio, 6))
        ws.cell(curr_r, c + 4, round(prev, 6) if prev is not None else "")
        diff = (ratio - prev) if prev is not None else ""
        ws.cell(curr_r, c + 5, round(diff, 6) if isinstance(diff, float) else "")
        for k in range(8):
            ws.cell(curr_r, c + k).border = BORDER

    ws.column_dimensions[get_column_letter(start_col - 1)].width = 3
    widths_s = [8, 12, 14, 12, 12, 12, 10, 10]
    for i, w in enumerate(widths_s):
        ws.column_dimensions[get_column_letter(c + i)].width = w


def fill_yj_summary_block_extras(ws, rank_df, start_col=22, row_4g=6, row_5g=7):
    if rank_df is None or rank_df.empty:
        return
    for sys_name, r in [("4G", row_4g), ("5G", row_5g)]:
        sub = rank_df[(rank_df["地市"] == "阳江") & (rank_df["制式"] == sys_name)]
        if len(sub) > 0:
            ws.cell(r, start_col + 6, int(sub["全省排名"].iloc[0]))
            ws.cell(r, start_col + 7, int(sub["三类排名"].iloc[0]))


def write_yj_region_block(ws, yj_4g, yj_5g, start_col=22, start_row=15):
    c = start_col
    areas = ["江城", "阳东", "阳西", "阳春", "南区"]
    band_cols = [
        ("4G", "FDD1800"), ("4G", "FDD900"), ("4G", "F频"), ("4G", "D频"), ("4G", "E频"), ("4G", "A频段"),
        ("5G", "700M"), ("5G", "2.6G"), ("5G", "4.9G"),
    ]

    def write_simple_table(ws, title, s_row):
        r = s_row
        if title:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + len(band_cols))
            style_header(ws.cell(r, c, title))
            r += 1
        hdr_row = r
        ws.cell(r, c, "区域")
        style_sub_header(ws.cell(r, c))
        for j, (_, b) in enumerate(band_cols):
            style_sub_header(ws.cell(r, c + 1 + j, b))
        style_sub_header(ws.cell(r, c + 1 + len(band_cols), "总计"))
        r += 1
        first_data_row = r
        for a in areas:
            ws.cell(r, c, a).border = BORDER
            area_col_letter = get_column_letter(c)
            for j, bk in enumerate(band_cols):
                band_col_letter = get_column_letter(c + 1 + j)
                if title:
                    formula = f'=COUNTIFS($B:$B,${area_col_letter}{r},$E:$E,{band_col_letter}${hdr_row},$H:$H,">-100")'
                else:
                    formula = f'=COUNTIFS($B:$B,${area_col_letter}{r},$E:$E,{band_col_letter}${hdr_row})'
                ws.cell(r, c + 1 + j, formula).border = BORDER
            start_b_letter = get_column_letter(c + 1)
            end_b_letter = get_column_letter(c + len(band_cols))
            ws.cell(r, c + 1 + len(band_cols), f'=SUM({start_b_letter}{r}:{end_b_letter}{r})').border = BORDER
            r += 1
        last_data_row = r - 1
        ws.cell(r, c, "总计").border = BORDER
        for j, bk in enumerate(band_cols):
            cur_col_letter = get_column_letter(c + 1 + j)
            ws.cell(r, c + 1 + j, f'=SUM({cur_col_letter}{first_data_row}:{cur_col_letter}{last_data_row})').border = BORDER
        tot_col_letter = get_column_letter(c + 1 + len(band_cols))
        ws.cell(r, c + 1 + len(band_cols), f'=SUM({tot_col_letter}{first_data_row}:{tot_col_letter}{last_data_row})').border = BORDER
        return (hdr_row, first_data_row, last_data_row, r)

    def write_combined_table(ws, s_row, t1_info, t2_info):
        _, t1_first, t1_last, t1_tot = t1_info
        _, t2_first, t2_last, t2_tot = t2_info
        r = s_row
        ws.cell(r, c, "区域")
        style_sub_header(ws.cell(r, c))
        for j, (_, b) in enumerate(band_cols):
            style_sub_header(ws.cell(r, c + 1 + j, b))
        style_sub_header(ws.cell(r, c + 1 + len(band_cols), "总计"))
        r += 1
        for idx, a in enumerate(areas):
            ws.cell(r, c, a).border = BORDER
            t1_r = t1_first + idx
            t2_r = t2_first + idx
            for j, bk in enumerate(band_cols):
                col_let = get_column_letter(c + 1 + j)
                ws.cell(r, c + 1 + j, f'={col_let}{t1_r}&"("&{col_let}{t2_r}&")"').border = BORDER
            tot_col_letter = get_column_letter(c + 1 + len(band_cols))
            ws.cell(r, c + 1 + len(band_cols), f'={tot_col_letter}{t1_r}&"("&{tot_col_letter}{t2_r}&")"').border = BORDER
            r += 1
        ws.cell(r, c, "总计").border = BORDER
        for j, bk in enumerate(band_cols):
            col_let = get_column_letter(c + 1 + j)
            ws.cell(r, c + 1 + j, f'={col_let}{t1_tot}&"("&{col_let}{t2_tot}&")"').border = BORDER
        tot_col_letter = get_column_letter(c + 1 + len(band_cols))
        ws.cell(r, c + 1 + len(band_cols), f'={tot_col_letter}{t1_tot}&"("&{tot_col_letter}{t2_tot}&")"').border = BORDER
        return r

    t1_info = write_simple_table(ws, None, start_row)
    t2_info = write_simple_table(ws, "干扰大于-100", t1_info[3] + 2)
    write_combined_table(ws, t2_info[3] + 2, t1_info, t2_info)

    for j in range(len(band_cols) + 1):
        ws.column_dimensions[get_column_letter(c + j)].width = 11


def write_province_ranking_block(ws, rank_df, start_col=22, start_row=57):
    c = start_col
    headers = ["地市", "4G总小区数", "4G干扰小区数", "4G干扰占比", "全省排名", "三类排名",
               "5G总小区数", "5G干扰小区数", "5G干扰占比", "全省排名", "三类排名"]
    for i, h in enumerate(headers):
        style_header(ws.cell(start_row, c + i, h))

    r = start_row + 1
    for cls in ["一类", "二类", "三类"]:
        for ci, city in enumerate(CITY_CLASS[cls]):
            ws.cell(r, c - 1, cls if ci == 0 else "").border = BORDER
            ws.cell(r, c, city).border = BORDER
            for s_i, system in enumerate(["4G", "5G"]):
                sub = rank_df[(rank_df["地市"] == city) & (rank_df["制式"] == system)]
                base = c + 1 + s_i * 5
                if len(sub) == 0:
                    for k in range(5):
                        ws.cell(r, base + k, 0).border = BORDER
                else:
                    ws.cell(r, base,     int(sub["总小区数"].iloc[0])).border = BORDER
                    ws.cell(r, base + 1, int(sub["干扰小区数"].iloc[0])).border = BORDER
                    ws.cell(r, base + 2, round(float(sub["干扰占比"].iloc[0]), 6)).border = BORDER
                    ws.cell(r, base + 3, int(sub["全省排名"].iloc[0])).border = BORDER
                    ws.cell(r, base + 4, int(sub["三类排名"].iloc[0])).border = BORDER
            r += 1

    widths_r = [10, 14, 14, 12, 10, 10, 14, 14, 12, 10, 10]
    for i, w in enumerate(widths_r):
        ws.column_dimensions[get_column_letter(c + i)].width = w


def write_full_sheet(ws, sheet_name, yj_4g, yj_5g, rank_data, day_cols,
                     yj_4g_total, yj_5g_total, prev_4g=None, prev_5g=None):
    ws.title = sheet_name
    write_main_table(ws, yj_4g, yj_5g, day_cols)
    rank_df = compute_province_ranking(rank_data)
    write_yj_summary_block(ws, yj_4g, yj_5g, yj_4g_total, yj_5g_total, prev_4g, prev_5g, start_col=22, start_row=4)
    fill_yj_summary_block_extras(ws, rank_df, start_col=22, row_4g=6, row_5g=7)
    write_yj_region_block(ws, yj_4g, yj_5g, start_col=22, start_row=15)
    write_province_ranking_block(ws, rank_df, start_col=22, start_row=57)


def generate_weekly_report(start_date_str: str, end_date_str: str, prev_report_path: str = None, log_fn=None) -> str:
    """生成周报主入口函数"""
    d_start = datetime.strptime(start_date_str, "%Y-%m-%d")
    d_end = datetime.strptime(end_date_str, "%Y-%m-%d")
    day_dates = [d_start + timedelta(days=i) for i in range((d_end - d_start).days + 1)]
    day_cols = [d.strftime("%m-%d") for d in day_dates]
    week_label = f"{d_start.strftime('%m%d')}-{d_end.strftime('%m%d')}"

    if log_fn:
        log_fn(f"开始生成周报: {week_label} (共 {len(day_dates)} 天)")

    rank_data_level = []
    rank_data_flag = []
    yj_records_level = []
    yj_records_flag = []

    # 4G 数据处理
    if log_fn:
        log_fn("正在处理全省 4G 干扰数据...")
    for city in ALL_CITIES:
        total, level_records, flag_records = process_city_4g(city, day_dates, log_fn)
        rank_data_level.append({"制式": "4G", "地市": city, "总小区数": total, "干扰小区数": len(level_records)})
        rank_data_flag.append({"制式": "4G", "地市": city, "总小区数": total, "干扰小区数": len(flag_records)})
        if city == "阳江":
            yj_records_level.extend(level_records)
            yj_records_flag.extend(flag_records)
        gc.collect()

    # 5G 数据处理
    if log_fn:
        log_fn("正在处理全省 5G 干扰数据...")
    for city in ALL_CITIES:
        total, level_records, flag_records = process_city_5g(city, day_dates, log_fn)
        rank_data_level.append({"制式": "5G", "地市": city, "总小区数": total, "干扰小区数": len(level_records)})
        rank_data_flag.append({"制式": "5G", "地市": city, "总小区数": total, "干扰小区数": len(flag_records)})
        if city == "阳江":
            yj_records_level.extend(level_records)
            yj_records_flag.extend(flag_records)
        gc.collect()

    yj_4g_level = filter_yangjiang([r for r in yj_records_level if r["制式"] == "4G"], "4G", log_fn)
    yj_5g_level = filter_yangjiang([r for r in yj_records_level if r["制式"] == "5G"], "5G", log_fn)
    yj_4g_flag = filter_yangjiang([r for r in yj_records_flag if r["制式"] == "4G"], "4G", log_fn)
    yj_5g_flag = filter_yangjiang([r for r in yj_records_flag if r["制式"] == "5G"], "5G", log_fn)

    yj_4g_total = next((r["总小区数"] for r in rank_data_level if r["制式"] == "4G" and r["地市"] == "阳江"), 0)
    yj_5g_total = next((r["总小区数"] for r in rank_data_level if r["制式"] == "5G" and r["地市"] == "阳江"), 0)

    # 历史数据对比读取
    prev_4g, prev_5g = None, None
    if prev_report_path and os.path.exists(prev_report_path):
        try:
            from openpyxl import load_workbook
            wb_p = load_workbook(prev_report_path, data_only=True)
            ws_p = wb_p[wb_p.sheetnames[0]]
            v4 = ws_p.cell(6, 25).value
            v5 = ws_p.cell(7, 25).value
            wb_p.close()
            prev_4g = float(v4) if isinstance(v4, (int, float)) else None
            prev_5g = float(v5) if isinstance(v5, (int, float)) else None
        except Exception as e:
            if log_fn:
                log_fn(f"[WARN] 读取历史周报失败: {e}")

    out_file = os.path.join(REPORT_DIR, f"全省的排名及干扰小区（{week_label}）_融合版.xlsx")
    wb = Workbook()

    # Sheet 1
    ws1 = wb.active
    write_full_sheet(ws1, "干扰电平判断", yj_4g_level, yj_5g_level, rank_data_level,
                     day_cols, yj_4g_total, yj_5g_total, prev_4g, prev_5g)

    # Sheet 2
    ws2 = wb.create_sheet()
    write_full_sheet(ws2, "干扰小区标识判断", yj_4g_flag, yj_5g_flag, rank_data_flag,
                     day_cols, yj_4g_total, yj_5g_total, prev_4g, prev_5g)

    # 自动归档至数据库
    try:
        from core.database import save_weekly_province_ranking, save_weekly_yj_cells, init_db
        init_db()
        save_weekly_province_ranking(start_date_str, end_date_str, "干扰电平判断", compute_province_ranking(rank_data_level))
        save_weekly_province_ranking(start_date_str, end_date_str, "干扰小区标识判断", compute_province_ranking(rank_data_flag))
        save_weekly_yj_cells(start_date_str, end_date_str, "干扰电平判断", yj_4g_level, yj_5g_level)
        save_weekly_yj_cells(start_date_str, end_date_str, "干扰小区标识判断", yj_4g_flag, yj_5g_flag)
        if log_fn:
            log_fn("周报数据已自动同步归档至 SQLite 数据库")
    except Exception as e:
        if log_fn:
            log_fn(f"[WARN] 数据库归档异常: {e}")

    wb.save(out_file)
    if log_fn:
        log_fn(f"周报生成成功: {out_file}")
    return out_file
