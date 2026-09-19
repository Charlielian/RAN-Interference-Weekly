"""
全省干扰排名及阳江干扰小区 周报融合脚本
========================================
功能：
  1) 读取 4G/5G 干扰小区源文件（按城市+日期拆分），流式处理；
  2) 同时按两种判断标准筛选干扰小区，输出一个Excel两个Sheet：
     Sheet1 - 干扰电平判断：按平均干扰电平门限筛选
             4G：> -105 dBm，5G-700M：> -105 dBm，5G-非700M：> -107 dBm，7天 >= 3天
     Sheet2 - 干扰小区标识判断：按源文件【是否干扰小区】字段 == TRUE，7天 >= 3天
  3) 每个Sheet保持与原报表一致的输出格式。

使用方法：
  1) python3 干扰小区周报_脚本_融合版.py
  2) 输出文件：全省的排名及干扰小区（{MMDD}-{MMDD}）_融合版.xlsx
"""

import os
import re
import gc
from datetime import datetime, timedelta
from collections import defaultdict

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIG
# ============================================================
# 源数据目录
BASE_DIR = "全省干扰排名及阳江的干扰小区清单"
DIR_4G   = os.path.join(BASE_DIR, "4G干扰文件")
DIR_5G   = os.path.join(BASE_DIR, "5G干扰文件")


def auto_detect_week():
    """自动检测数据日期范围"""
    import glob
    files = glob.glob(os.path.join(DIR_4G, "4G干扰小区_*.xlsx"))
    dates = set()
    for f in files:
        m = re.search(r"(\d{4}-\d{2}-\d{2})", os.path.basename(f))
        if m:
            dates.add(datetime.strptime(m.group(1), "%Y-%m-%d"))
    if not dates:
        return datetime(2026, 6, 15), datetime(2026, 6, 21)
    dates = sorted(dates)
    return dates[0], dates[-1]

# 自动检测本周起止日期
WEEK_START, WEEK_END = auto_detect_week()
WEEK_LABEL  = f"{WEEK_START.strftime('%m%d')}-{WEEK_END.strftime('%m%d')}"
OUTPUT_PATH = f"全省的排名及干扰小区（{WEEK_LABEL}）_融合版.xlsx"

# 全省 21 地市
CITY_CLASS = {
    "一类": ["广州", "深圳", "东莞", "佛山"],
    "二类": ["惠州", "珠海", "中山", "汕头", "湛江", "江门", "茂名", "揭阳"],
    "三类": ["清远", "肇庆", "梅州", "韶关", "河源", "潮州", "阳江", "汕尾", "云浮"],
}
ALL_CITIES = [c for v in CITY_CLASS.values() for c in v]

# 干扰门限 (dBm) —— 用于干扰电平判断
THRESHOLD_4G      = -105
THRESHOLD_5G_700  = -105
THRESHOLD_5G_OTH  = -107

# 5G 频段映射（源文件"频段"列 -> 报告内显示名）
BAND_5G_MAP = {
    "700M":   "700M",
    "2.6GHz": "2.6G",
    "4.9GHz": "4.9G",
}

# 4G 频段映射
BAND_4G_MAP = {
    "FD":   "FDD1800",
    "FG":   "FDD900",
    "LF":   "F频",
    "LD":   "D频",
    "LE":   "E频",
    "A频段": "A频段",
    "NB":   "NB-IoT",
}

# 阳江区域提取
YJ_AREA_RULES = [
    (re.compile(r"^阳江阳西"), "阳西"),
    (re.compile(r"^阳江江城"), "江城"),
    (re.compile(r"^阳江南区"), "南区"),
    (re.compile(r"^阳江阳东"), "阳东"),
    (re.compile(r"^阳江阳春"), "阳春"),
]

# 历史xlsx路径（用于对比上周）
PREVIOUS_REPORT_PATH = ""

# 处理日志文件
LOG_PATH = f"处理日志_{WEEK_LABEL}_融合版.log"


# ============================================================
# 工具函数
# ============================================================
def extract_yj_area(cell_name):
    if not isinstance(cell_name, str):
        return ""
    for pat, area in YJ_AREA_RULES:
        if pat.match(cell_name):
            return area
    return ""


_log_file = None


def log_message(msg: str):
    """同时输出到控制台和日志文件"""
    print(msg)
    global _log_file
    if _log_file is None:
        _log_file = open(LOG_PATH, "a", encoding="utf-8")
    _log_file.write(msg + "\n")
    _log_file.flush()


def close_log():
    """关闭日志文件"""
    global _log_file
    if _log_file is not None:
        _log_file.close()
        _log_file = None


def map_band_4g(b):
    return BAND_4G_MAP.get(b, b or "")


def map_band_5g(b):
    return BAND_5G_MAP.get(b, b or "")


def threshold_for(system, band):
    if system == "4G":
        return THRESHOLD_4G
    return THRESHOLD_5G_700 if band == "700M" else THRESHOLD_5G_OTH


# ============================================================
# 流式处理单个地市数据（同时返回两种判断结果）
# ============================================================
def process_city_4g(city: str, day_dates: list) -> tuple:
    """
    处理单个地市4G数据，一次性读取并同时按两种标准判断。
    返回 (total_cgi, level_records, flag_records)
      level_records: 按平均干扰电平门限判断
      flag_records:  按【是否干扰小区】字段判断
    """
    all_data = []

    for date in day_dates:
        file_path = os.path.join(DIR_4G, f"4G干扰小区_{date.strftime('%Y-%m-%d')}_{city}.xlsx")
        if not os.path.exists(file_path):
            continue

        try:
            df = pd.read_excel(file_path, dtype={"CGI": str})
        except Exception as e:
            log_message(f"    [WARN] 读取失败，跳过: {file_path} ({type(e).__name__}: {e})")
            continue
        df = df[["CGI", "小区名", "频段", "平均干扰电平", "是否干扰小区"]].copy()
        df["日期"] = date
        df["CGI"] = df["CGI"].astype(str).str.strip()
        df = df[df["CGI"] != "None"]
        all_data.append(df)
        del df
        gc.collect()

    if not all_data:
        return 0, [], []

    combined = pd.concat(all_data, ignore_index=True)
    del all_data
    gc.collect()

    agg = combined.groupby(["CGI", "日期"]).agg({
        "小区名": "first",
        "频段": "first",
        "平均干扰电平": "mean",
        "是否干扰小区": "first"
    }).reset_index()

    pivot = agg.pivot(index="CGI", columns="日期", values="平均干扰电平")
    pivot_interf = agg.pivot(index="CGI", columns="日期", values="是否干扰小区")
    meta = agg.groupby("CGI").agg({"小区名": "first", "频段": "first"})

    del combined, agg
    gc.collect()

    total = len(pivot)
    level_records = []   # 按干扰电平判断
    flag_records = []    # 按干扰小区标识判断

    for cgi in pivot.index:
        row_meta = meta.loc[cgi]
        band = str(row_meta["频段"]) if row_meta["频段"] else ""
        day_vals = {d: pivot.loc[cgi, d] if d in pivot.columns else None
                    for d in day_dates}
        day_interf = {d: pivot_interf.loc[cgi, d] if d in pivot_interf.columns else None
                      for d in day_dates}
        valid = [v for v in day_vals.values() if v is not None and not np.isnan(v)]
        avg = sum(valid) / len(valid) if valid else float("nan")
        thr = THRESHOLD_4G

        # --- 判断1：按干扰电平门限 ---
        level_exceed_days = sum(1 for v in valid if v > thr)
        if level_exceed_days >= 3:
            record = {
                "制式": "4G",
                "地市": city,
                "CGI": cgi,
                "小区名": row_meta["小区名"],
                "频段": band,
                "门限": thr,
                "干扰天数": level_exceed_days,
                "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            level_records.append(record)

        # --- 判断2：按是否干扰小区标识 ---
        flag_exceed_days = sum(1 for v in day_interf.values() if v is not None and v == True)
        if flag_exceed_days >= 3:
            record = {
                "制式": "4G",
                "地市": city,
                "CGI": cgi,
                "小区名": row_meta["小区名"],
                "频段": band,
                "门限": thr,
                "干扰天数": flag_exceed_days,
                "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            flag_records.append(record)

    del pivot, pivot_interf, meta
    gc.collect()
    return total, level_records, flag_records


def process_city_5g(city: str, day_dates: list) -> tuple:
    """
    处理单个地市5G数据，一次性读取并同时按两种标准判断。
    返回 (total_cgi, level_records, flag_records)
    """
    all_data = []

    for date in day_dates:
        file_path = os.path.join(DIR_5G, f"5G干扰小区_{date.strftime('%Y-%m-%d')}_{city}.xlsx")
        if not os.path.exists(file_path):
            continue

        try:
            df = pd.read_excel(file_path, dtype={"CGI": str})
        except Exception as e:
            log_message(f"    [WARN] 读取失败，跳过: {file_path} ({type(e).__name__}: {e})")
            continue
        df = df[["CGI", "小区名", "频段", "全频段均值", "是否干扰小区"]].copy()
        df["日期"] = date
        df["CGI"] = df["CGI"].astype(str).str.strip()
        df = df[df["CGI"] != "None"]
        all_data.append(df)
        del df
        gc.collect()

    if not all_data:
        return 0, [], []

    combined = pd.concat(all_data, ignore_index=True)
    del all_data
    gc.collect()

    agg = combined.groupby(["CGI", "日期"]).agg({
        "小区名": "first",
        "频段": "first",
        "全频段均值": "mean",
        "是否干扰小区": "first"
    }).reset_index()

    pivot = agg.pivot(index="CGI", columns="日期", values="全频段均值")
    pivot_interf = agg.pivot(index="CGI", columns="日期", values="是否干扰小区")
    meta = agg.groupby("CGI").agg({"小区名": "first", "频段": "first"})

    del combined, agg
    gc.collect()

    total = len(pivot)
    level_records = []
    flag_records = []

    for cgi in pivot.index:
        row_meta = meta.loc[cgi]
        band = str(row_meta["频段"]) if row_meta["频段"] else ""
        thr = THRESHOLD_5G_700 if "700" in band.upper() else THRESHOLD_5G_OTH
        day_vals = {d: pivot.loc[cgi, d] if d in pivot.columns else None
                    for d in day_dates}
        day_interf = {d: pivot_interf.loc[cgi, d] if d in pivot_interf.columns else None
                      for d in day_dates}
        valid = [v for v in day_vals.values() if v is not None and not np.isnan(v)]
        avg = sum(valid) / len(valid) if valid else float("nan")

        # --- 判断1：按干扰电平门限 ---
        level_exceed_days = sum(1 for v in valid if v > thr)
        if level_exceed_days >= 3:
            record = {
                "制式": "5G",
                "地市": city,
                "CGI": cgi,
                "小区名": row_meta["小区名"],
                "频段": band,
                "门限": thr,
                "干扰天数": level_exceed_days,
                "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            level_records.append(record)

        # --- 判断2：按是否干扰小区标识 ---
        flag_exceed_days = sum(1 for v in day_interf.values() if v is not None and v == True)
        if flag_exceed_days >= 3:
            record = {
                "制式": "5G",
                "地市": city,
                "CGI": cgi,
                "小区名": row_meta["小区名"],
                "频段": band,
                "门限": thr,
                "干扰天数": flag_exceed_days,
                "平均干扰电平": avg,
            }
            for d in day_dates:
                record[d.strftime("%m-%d")] = day_vals.get(d)
            flag_records.append(record)

    del pivot, pivot_interf, meta
    gc.collect()
    return total, level_records, flag_records


# ============================================================
# 数据筛选与处理
# ============================================================
def filter_yangjiang(records: list, system: str) -> pd.DataFrame:
    """筛选阳江数据并处理"""
    df = pd.DataFrame(records)
    if df.empty:
        return df

    df = df[df["地市"] == "阳江"].copy()
    df["区域"] = df["小区名"].apply(extract_yj_area)
    unknown = df[df["区域"] == ""]
    if len(unknown) > 0:
        log_message(f"  [WARN] {system} 以下小区无法识别区域:")
        for _, row in unknown.iterrows():
            log_message(f"         {row['小区名']}")
        log_message(f"  [WARN] {system} 无法识别区域的小区数: {len(unknown)}")

    if system == "4G":
        df["频段"] = df["频段"].apply(map_band_4g)
    else:
        df["频段"] = df["频段"].apply(map_band_5g)

    df = df.sort_values(by=["干扰天数", "平均干扰电平"],
                        ascending=[False, False]).reset_index(drop=True)
    return df


def compute_province_ranking(rank_data: list) -> pd.DataFrame:
    """计算全省排名"""
    df = pd.DataFrame(rank_data)
    if df.empty:
        return pd.DataFrame()

    df["干扰占比"] = df.apply(
        lambda r: r["干扰小区数"] / r["总小区数"] if r["总小区数"] else 0.0,
        axis=1)
    df["全省排名"] = df.groupby("制式")["干扰占比"].rank(
        method="min", ascending=True).astype(int)
    df["类别"] = df["地市"].map(
        {c: k for k, v in CITY_CLASS.items() for c in v})
    df["三类排名"] = df.groupby(["制式", "类别"])["干扰占比"].rank(
        method="min", ascending=True).astype(int)
    return df[["类别", "地市", "制式", "总小区数", "干扰小区数",
               "干扰占比", "全省排名", "三类排名"]]


def try_load_prev_ratios(path: str):
    """读取历史报表的干扰占比"""
    if not path or not os.path.exists(path):
        return None, None
    try:
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        ws = wb[wb.sheetnames[0]]
        v4g = ws.cell(6, 24).value
        v5g = ws.cell(7, 24).value
        wb.close()
        return (float(v4g) if isinstance(v4g, (int, float)) else None,
                float(v5g) if isinstance(v5g, (int, float)) else None)
    except Exception as e:
        print(f"[WARN] 读取历史失败: {e}")
        return None, None


# ============================================================
# 输出 Excel — 样式常量
# ============================================================
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL  = PatternFill(start_color="305496", end_color="305496", fill_type="solid")
SUB_FILL   = PatternFill(start_color="8FAADC", end_color="8FAADC", fill_type="solid")
HEAD_FONT  = Font(bold=True, color="FFFFFF")


def style_header(cell):
    cell.font = HEAD_FONT
    cell.fill = HEAD_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def style_sub_header(cell):
    cell.font = HEAD_FONT
    cell.fill = SUB_FILL
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


# ============================================================
# 输出 Excel — 各区块写入函数
# ============================================================
def write_main_table(ws, yj_4g, yj_5g, day_cols, start_row=1, start_col=1):
    """写入阳江干扰小区明细表"""
    headers = (["制式", "地市", "小区名", "CGI", "区域", "频段",
                "门限(dBm)", "干扰天数", "平均干扰电平(dBm)"]
               + [f"{dc}干扰值(dBm)" for dc in day_cols])
    for c, h in enumerate(headers, start_col):
        style_header(ws.cell(start_row, c, h))

    r = start_row + 1
    for system, df in (("4G", yj_4g), ("5G", yj_5g)):
        for _, row in df.iterrows():
            ws.cell(r, start_col,     system)
            ws.cell(r, start_col + 1, "阳江")
            ws.cell(r, start_col + 2, row.get("小区名", ""))
            ws.cell(r, start_col + 3, row.get("CGI", ""))
            ws.cell(r, start_col + 4, row.get("区域", ""))
            ws.cell(r, start_col + 5, row.get("频段", ""))
            ws.cell(r, start_col + 6, int(row.get("门限", 0)))
            ws.cell(r, start_col + 7, int(row.get("干扰天数", 0)))
            v = row.get("平均干扰电平", 0)
            ws.cell(r, start_col + 8, round(float(v), 2) if pd.notna(v) else "")
            for i, dc in enumerate(day_cols):
                v = row.get(dc)
                if pd.notna(v):
                    ws.cell(r, start_col + 9 + i, round(float(v), 5))
            for c in range(start_col, start_col + 9 + len(day_cols)):
                ws.cell(r, c).border = BORDER
            r += 1

    widths = [6, 8, 38, 22, 10, 8, 10, 10, 18] + [16] * len(day_cols)
    for i, w in enumerate(widths, start_col):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_yj_summary_block(ws, yj_4g, yj_5g, total_4g_yj, total_5g_yj,
                           prev_4g, prev_5g, start_col=22, start_row=4):
    """写入阳江干扰概况汇总"""
    cnt_4g = len(yj_4g)
    cnt_5g = len(yj_5g)
    ratio_4g = (cnt_4g / total_4g_yj) if total_4g_yj else 0
    ratio_5g = (cnt_5g / total_5g_yj) if total_5g_yj else 0
    c = start_col

    title = ws.cell(start_row, c, f"45G干扰情况（{WEEK_LABEL}）")
    title.font = Font(bold=True, color="FFFFFF")
    title.fill = HEAD_FILL
    title.alignment = Alignment(horizontal="center", vertical="center")
    title.border = BORDER
    ws.merge_cells(start_row=start_row, start_column=c,
                   end_row=start_row, end_column=c + 8)
    for col in range(c, c + 9):
        ws.cell(start_row, col).border = BORDER

    sub_headers = ["制式", "高干扰小区（括号内为大于-100dBm的数量）",
                   "干扰比例", "对比上周", "全省排名",
                   "全省干扰小区比例平均值", "三类排名",
                   "三类干扰小区比例平均值"]
    for i, h in enumerate(sub_headers):
        style_sub_header(ws.cell(start_row + 1, c + i, h))

    for s_i, (system, cnt, ratio, prev) in enumerate([
        ("4G", cnt_4g, ratio_4g, prev_4g),
        ("5G", cnt_5g, ratio_5g, prev_5g),
    ]):
        r = start_row + 2 + s_i
        ws.cell(r, c,     system)
        ws.cell(r, c + 1, f"{cnt}")
        ws.cell(r, c + 2, round(ratio, 6))
        if prev is not None:
            ws.cell(r, c + 3, round(ratio - prev, 6))
        for col in range(c, c + 4):
            ws.cell(r, col).border = BORDER


def fill_yj_summary_block_extras(ws, rank_df, start_col=22, row_4g=6, row_5g=7):
    """填写阳江概况汇总中的排名等额外信息"""
    c = start_col
    full_avg, cat_avg = {}, {}
    for s in ("4G", "5G"):
        sub = rank_df[rank_df["制式"] == s]
        full_avg[s] = sub["干扰占比"].mean() if len(sub) else 0
        cat_avg[s] = (sub.groupby("类别")["干扰占比"].mean().to_dict()
                      if len(sub) else {})
    for s_i, system in enumerate(["4G", "5G"]):
        r = row_4g + s_i
        sub = rank_df[(rank_df["地市"] == "阳江") & (rank_df["制式"] == system)]
        if len(sub) == 0:
            continue
        ws.cell(r, c + 4, int(sub["全省排名"].iloc[0]))
        ws.cell(r, c + 5, round(full_avg[system], 6))
        ws.cell(r, c + 6, int(sub["三类排名"].iloc[0]))
        ws.cell(r, c + 7, round(cat_avg[system].get("三类", 0), 6))
        for col in range(c + 4, c + 8):
            ws.cell(r, col).border = BORDER


def write_yj_region_block(ws, yj_4g, yj_5g, start_col=22, start_row=15):
    """写入阳江各区域频段分布"""
    c = start_col
    band_keys = ["FDD1800", "FDD900", "F频", "D频", "E频", "A频段", "700M", "2.6G", "4.9G"]
    band_cols = band_keys + ["总计"]
    areas = ["江城", "阳东", "阳西", "阳春", "南区"]

    def write_simple_table(ws, title, start_row):
        r = start_row
        if title:
            ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + len(band_keys) + 1)
            style_header(ws.cell(r, c, title))
            r += 1
        hdr_row = r
        for i, h in enumerate(["区域"] + band_cols):
            style_sub_header(ws.cell(r, c + i, h))
        r += 1
        first_data_row = r
        area_col_letter = get_column_letter(c)
        for area in areas:
            ws.cell(r, c, area).border = BORDER
            for j, b in enumerate(band_keys):
                b_col_letter = get_column_letter(c + 1 + j)
                if title:
                    formula = f'=COUNTIFS($B:$B,${area_col_letter}{r},$E:$E,{b_col_letter}${hdr_row},$H:$H,">-100")'
                else:
                    formula = f'=COUNTIFS($B:$B,${area_col_letter}{r},$E:$E,{b_col_letter}${hdr_row})'
                ws.cell(r, c + 1 + j, formula).border = BORDER
            start_b_letter = get_column_letter(c + 1)
            end_b_letter = get_column_letter(c + len(band_keys))
            ws.cell(r, c + 1 + len(band_keys), f'=SUM({start_b_letter}{r}:{end_b_letter}{r})').border = BORDER
            r += 1
        last_data_row = r - 1
        ws.cell(r, c, "总计").border = BORDER
        for j in range(len(band_keys)):
            col_letter = get_column_letter(c + 1 + j)
            ws.cell(r, c + 1 + j, f'=SUM({col_letter}{first_data_row}:{col_letter}{last_data_row})').border = BORDER
        tot_col_letter = get_column_letter(c + 1 + len(band_keys))
        ws.cell(r, c + 1 + len(band_keys), f'=SUM({tot_col_letter}{first_data_row}:{tot_col_letter}{last_data_row})').border = BORDER
        return (hdr_row, first_data_row, last_data_row, r)

    def write_combined_table(ws, start_row, t1_info, t2_info):
        _, t1_first, t1_last, t1_tot = t1_info
        _, t2_first, t2_last, t2_tot = t2_info
        r = start_row
        for i, h in enumerate(["区域"] + band_cols):
            style_sub_header(ws.cell(r, c + i, h))
        r += 1
        for idx, area in enumerate(areas):
            ws.cell(r, c, area).border = BORDER
            t1_r = t1_first + idx
            t2_r = t2_first + idx
            for j in range(len(band_keys)):
                col_let = get_column_letter(c + 1 + j)
                ws.cell(r, c + 1 + j, f'={col_let}{t1_r}&"("&{col_let}{t2_r}&")"').border = BORDER
            tot_col = get_column_letter(c + 1 + len(band_keys))
            ws.cell(r, c + 1 + len(band_keys), f'={tot_col}{t1_r}&"("&{tot_col}{t2_r}&")"').border = BORDER
            r += 1
        ws.cell(r, c, "总计").border = BORDER
        for j in range(len(band_keys)):
            col_let = get_column_letter(c + 1 + j)
            ws.cell(r, c + 1 + j, f'={col_let}{t1_tot}&"("&{col_let}{t2_tot}&")"').border = BORDER
        tot_col = get_column_letter(c + 1 + len(band_keys))
        ws.cell(r, c + 1 + len(band_keys), f'={tot_col}{t1_tot}&"("&{tot_col}{t2_tot}&")"').border = BORDER
        return r

    t1_info = write_simple_table(ws, None, start_row)
    t2_info = write_simple_table(ws, "干扰大于-100", t1_info[3] + 2)
    write_combined_table(ws, t2_info[3] + 2, t1_info, t2_info)

    for j in range(len(band_cols) + 1):
        ws.column_dimensions[get_column_letter(c + j)].width = 11


def write_province_ranking_block(ws, rank_df, start_col=22, start_row=57):
    """写入全省21地市排名"""
    c = start_col
    headers = ["地市", "4G总小区数", "4G干扰小区数", "4G干扰占比",
               "全省排名", "三类排名",
               "5G总小区数", "5G干扰小区数", "5G干扰占比",
               "全省排名", "三类排名"]
    for i, h in enumerate(headers):
        style_header(ws.cell(start_row, c + i, h))

    r = start_row + 1
    for cls in ["一类", "二类", "三类"]:
        for ci, city in enumerate(CITY_CLASS[cls]):
            ws.cell(r, c - 1, cls if ci == 0 else "").border = BORDER
            ws.cell(r, c, city).border = BORDER
            for s_i, system in enumerate(["4G", "5G"]):
                sub = rank_df[(rank_df["地市"] == city) & (rank_df["制式"] == system)]
                base = c + 1 + s_i * 5
                if len(sub) == 0:
                    for k in range(5):
                        ws.cell(r, base + k, 0)
                else:
                    ws.cell(r, base,     int(sub["总小区数"].iloc[0]))
                    ws.cell(r, base + 1, int(sub["干扰小区数"].iloc[0]))
                    ws.cell(r, base + 2, round(float(sub["干扰占比"].iloc[0]), 6))
                    ws.cell(r, base + 3, int(sub["全省排名"].iloc[0]))
                    ws.cell(r, base + 4, int(sub["三类排名"].iloc[0]))
                for k in range(5):
                    ws.cell(r, base + k).border = BORDER
            r += 1

    widths_r = [10, 14, 14, 12, 10, 10, 14, 14, 12, 10, 10]
    for i, w in enumerate(widths_r):
        ws.column_dimensions[get_column_letter(c + i)].width = w


# ============================================================
# 写入单个完整 Sheet
# ============================================================
def write_sheet(ws, sheet_name, yj_4g, yj_5g, rank_data, day_cols,
                yj_4g_total, yj_5g_total, prev_4g, prev_5g):
    """写入一个完整Sheet的内容（明细表 + 汇总 + 区域分布 + 全省排名）"""
    ws.title = sheet_name
    write_main_table(ws, yj_4g, yj_5g, day_cols)
    rank_df = compute_province_ranking(rank_data)
    write_yj_summary_block(ws, yj_4g, yj_5g, yj_4g_total, yj_5g_total,
                           prev_4g, prev_5g, start_col=22, start_row=4)
    fill_yj_summary_block_extras(ws, rank_df, start_col=22, row_4g=6, row_5g=7)
    write_yj_region_block(ws, yj_4g, yj_5g, start_col=22, start_row=15)
    write_province_ranking_block(ws, rank_df, start_col=22, start_row=57)


# ============================================================
# 主流程
# ============================================================
def main():
    day_dates = [WEEK_START + timedelta(days=i) for i in range((WEEK_END - WEEK_START).days + 1)]
    day_cols = [d.strftime("%m-%d") for d in day_dates]
    log_message(f"[INFO] 处理周区间: {WEEK_LABEL} (共 {len(day_dates)} 天)")

    # 两个判断标准的汇总数据
    rank_data_level = []   # 按干扰电平
    rank_data_flag = []    # 按干扰小区标识
    yj_records_level = []
    yj_records_flag = []

    # 流式处理4G
    log_message("[INFO] 处理 4G 源文件 (流式) ...")
    for ci, city in enumerate(ALL_CITIES, 1):
        total, level_records, flag_records = process_city_4g(city, day_dates)
        rank_data_level.append({
            "制式": "4G",
            "地市": city,
            "总小区数": total,
            "干扰小区数": len(level_records),
        })
        rank_data_flag.append({
            "制式": "4G",
            "地市": city,
            "总小区数": total,
            "干扰小区数": len(flag_records),
        })
        if city == "阳江":
            yj_records_level.extend(level_records)
            yj_records_flag.extend(flag_records)
        log_message(f"    {city} done: total={total}, level={len(level_records)}, flag={len(flag_records)}")
        gc.collect()

    # 流式处理5G
    log_message("[INFO] 处理 5G 源文件 (流式) ...")
    for ci, city in enumerate(ALL_CITIES, 1):
        total, level_records, flag_records = process_city_5g(city, day_dates)
        rank_data_level.append({
            "制式": "5G",
            "地市": city,
            "总小区数": total,
            "干扰小区数": len(level_records),
        })
        rank_data_flag.append({
            "制式": "5G",
            "地市": city,
            "总小区数": total,
            "干扰小区数": len(flag_records),
        })
        if city == "阳江":
            yj_records_level.extend(level_records)
            yj_records_flag.extend(flag_records)
        log_message(f"    {city} done: total={total}, level={len(level_records)}, flag={len(flag_records)}")
        gc.collect()

    # 筛选阳江数据
    yj_4g_level = filter_yangjiang([r for r in yj_records_level if r["制式"] == "4G"], "4G")
    yj_5g_level = filter_yangjiang([r for r in yj_records_level if r["制式"] == "5G"], "5G")
    yj_4g_flag = filter_yangjiang([r for r in yj_records_flag if r["制式"] == "4G"], "4G")
    yj_5g_flag = filter_yangjiang([r for r in yj_records_flag if r["制式"] == "5G"], "5G")

    log_message(f"[INFO] 干扰电平判断: 阳江 4G={len(yj_4g_level)}, 5G={len(yj_5g_level)}")
    log_message(f"[INFO] 干扰小区标识判断: 阳江 4G={len(yj_4g_flag)}, 5G={len(yj_5g_flag)}")

    # 阳江总小区数（两种判断共用相同的小区总数）
    yj_4g_total = next((r["总小区数"] for r in rank_data_level if r["制式"] == "4G" and r["地市"] == "阳江"), 0)
    yj_5g_total = next((r["总小区数"] for r in rank_data_level if r["制式"] == "5G" and r["地市"] == "阳江"), 0)
    log_message(f"[INFO] 阳江 4G 总小区数: {yj_4g_total}, 5G 总小区数: {yj_5g_total}")

    # 历史对比
    prev_4g, prev_5g = try_load_prev_ratios(PREVIOUS_REPORT_PATH)
    if prev_4g is not None:
        log_message(f"[INFO] 上周 4G 干扰占比={prev_4g}, 5G 干扰占比={prev_5g}")

    # 输出 Excel（两个Sheet）
    wb = Workbook()

    # Sheet 1: 干扰电平判断
    ws_level = wb.active
    write_sheet(ws_level, "干扰电平判断",
                yj_4g_level, yj_5g_level, rank_data_level,
                day_cols, yj_4g_total, yj_5g_total, prev_4g, prev_5g)

    # Sheet 2: 干扰小区标识判断
    ws_flag = wb.create_sheet()
    write_sheet(ws_flag, "干扰小区标识判断",
                yj_4g_flag, yj_5g_flag, rank_data_flag,
                day_cols, yj_4g_total, yj_5g_total, prev_4g, prev_5g)

    wb.save(OUTPUT_PATH)
    log_message(f"[OK] 已生成: {OUTPUT_PATH}")
    log_message("  Sheet1: 干扰电平判断")
    log_message("  Sheet2: 干扰小区标识判断")
    log_message("[INFO] 处理完成")
    close_log()


if __name__ == "__main__":
    main()